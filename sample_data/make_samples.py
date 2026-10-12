"""
Generate the demo files in this folder.

    python sample_data/make_samples.py

  • invoices.xlsx       – an invoice register (one row per invoice, header
                          buried under a report title)
  • Invoice_0761.xlsx   – a single printed invoice "form" (label/value cells)
  • ledger.xlsx         – an ERP-style subledger export: metadata rows above
                          the header, vendor context lines, opening/closing
                          balance rows and invoice references hidden in the
                          narration text
"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

OUT = Path(__file__).resolve().parent


def invoices_register() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    ws.append(["Purchase Invoice Register"])
    ws.append(["Period: February 2026"])
    ws.append([])
    ws.append(["Invoice No.", "Vendor", "Invoice Date", "Amount (NPR)"])
    rows = [
        ["INV-0758",    "Himalayan Traders Pvt. Ltd.",  "02/02/2026", "45,000.00"],
        ["INV-0759",    "Kathmandu Supplies Co.",       "05/02/2026", "12,500.50"],
        ["1043",        "Nepal Electronics Ltd",        "08/02/2026", "7,800.00"],
        ["EG/2026/118", "Everest Goods Inc",            "10/02/2026", "33,200.00"],
        ["INV-0762",    "Annapurna Paper Mills",        "12/02/2026", "18,450.00"],
        ["SB-2290",     "Sagarmatha Builders",          "14/02/2026", "96,000.00"],
        ["GL-5521",     "Gorkha Logistics",             "18/02/2026", "5,400.00"],
    ]
    for r in rows:
        ws.append(r)
    ws.append([])
    ws.append(["Grand Total", "", "", "2,18,350.50"])
    for cell in ws[4]:
        cell.font = Font(bold=True)
    ws["A1"].font = Font(bold=True, size=14)
    for col, width in zip("ABCD", (16, 32, 16, 16)):
        ws.column_dimensions[col].width = width
    wb.save(OUT / "invoices.xlsx")


def invoice_form() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws["B1"] = "Bhagwati Steel Industries Ltd."
    ws["B1"].font = Font(bold=True, size=14)
    ws["B2"] = "Hetauda Industrial Area, Makwanpur"
    ws["B4"] = "TAX INVOICE"
    ws["B4"].font = Font(bold=True)
    ws["A6"], ws["B6"] = "Invoice No.:", "0761"
    ws["D6"], ws["E6"] = "Invoice Date:", "22/02/2026"
    ws["A7"], ws["B7"] = "Customer:", "Sunrise Fabrication Pvt. Ltd."
    ws["D7"], ws["E7"] = "Invoice Miti:", "2082-11-10"
    ws.append([])
    ws.append(["S.N.", "Description", "Qty", "Rate", "Amount"])
    ws.append([1, "MS Angle 50x50x5", 40, 1500, 60000])
    ws.append([2, "MS Flat 40x6", 25, 1600, 40000])
    ws.append([])
    ws.append(["", "", "", "Taxable Amount:", 100000])
    ws.append(["", "", "", "VAT 13%:", 13000])
    ws.append(["", "", "", "Grand Total:", 113000])
    for col, width in zip("ABCDE", (14, 34, 14, 18, 14)):
        ws.column_dimensions[col].width = width
    wb.save(OUT / "Invoice_0761.xlsx")


def ledger() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Subledger"
    ws.append(["Sunrise Fabrication Pvt. Ltd."])
    ws.append(["Fiscal Year 2082/83"])
    ws.append(["Subledger Detail Report — Sundry Creditors"])
    ws.append([])
    ws.append(["Date", "Miti", "DOC No.", "GL - Ledger / SubLedger / Narration",
               "Debit Amount", "Credit Amount", "Balance Amount"])
    for cell in ws[5]:
        cell.font = Font(bold=True)

    def d(s: str) -> datetime:
        return datetime.strptime(s, "%Y-%m-%d")

    def vendor(code: str, name: str) -> None:
        ws.append([f"{code}  [ID : {name}]"])

    ws.append(["Subledger Opening B/L", "", "", "", "", "", "0.00"])

    vendor("S701", "HIMALAYAN TRADERS PVT. LTD.")
    ws.append([d("2026-02-02"), "2082-10-19", "PI/RM-L/82-83/0101\n(Post Journal)",
               "Purchase of raw material. Being Payable Against Ref. Doc No. INV-0758 Dated: 02-Feb-26",
               "", "45,000.00", "45,000.00"])

    vendor("S702", "KTM SUPPLIES COMPANY")
    ws.append([d("2026-02-06"), "2082-10-23", "PI/RM-L/82-83/0104",
               "Office consumables, Inv No. INV0759", "", "12,500.50", "57,500.50"])

    vendor("S703", "NEPAL ELECTRONICS")
    ws.append([d("2026-02-08"), "2082-10-25", "PI/RM-L/82-83/0107",
               "Electrical items Purchase Through Bill No 1043", "", "7,750.00", "65,250.50"])

    vendor("S704", "EVEREST GOODS INC")
    ws.append([d("2026-02-10"), "2082-10-27", "PI/RM-L/82-83/0110",
               "Packing material, Invoice No. EG/2026/118", "", "33,200.00", "98,450.50"])

    vendor("S705", "ANNAPURNA PAPER MILLS PVT. LTD.")
    ws.append([d("2026-02-15"), "2082-11-03", "PI/RM-L/82-83/0113",
               "Paper rolls, Ref. Doc No. INV-0726", "", "18,450.00", "1,16,900.50"])

    vendor("S706", "GORKHA LOGISTIC SERVICES")
    ws.append([d("2026-02-24"), "2082-11-12", "PI/RM-L/82-83/0118",
               "Freight charges, Inv No. GL5521", "", "5,400.00", "1,22,300.50"])

    vendor("S712", "BHAGWATI STEEL INDUSTRIES LTD.")
    ws.append([d("2026-02-22"), "2082-11-10", "PI/RM-L/82-83/0116",
               "MS Angle & Flat. Being Payable Against Ref. Doc No. 0761 Dated: 22-Feb-26 - Bhagwati Steels",
               "", "1,13,000.00", "2,35,300.50"])

    vendor("S715", "POKHARA FURNISHERS")
    ws.append([d("2026-02-14"), "2082-11-02", "PI/RM-L/82-83/0112",
               "Office furniture, Bill No 7781", "", "12,000.00", "2,47,300.50"])

    vendor("S720", "LUMBINI TRADING CONCERN")
    ws.append([d("2025-09-03"), "2082-05-18", "PI/RM-L/82-83/0042",
               "Old purchase carried forward, Bill No 3310", "", "8,900.00", "2,56,200.50"])

    ws.append(["Operation Total", "", "", "", "0.00", "2,56,200.50", ""])
    ws.append(["Closing Balance", "", "", "", "", "", "2,56,200.50"])
    ws.append([])
    ws.append(["Printed on 01-Mar-2026 by admin"])

    for col, width in zip("ABCDEFG", (22, 12, 26, 70, 14, 14, 16)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=6, max_col=1):
        if isinstance(row[0].value, datetime):
            row[0].number_format = "yyyy-mm-dd"
    wb.save(OUT / "ledger.xlsx")


if __name__ == "__main__":
    invoices_register()
    invoice_form()
    ledger()
    print(f"Sample files written to {OUT}")
